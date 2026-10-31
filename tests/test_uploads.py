"""Uploads: format validation, parsing of PDF / Word / Excel / TSV, and adding documents to a live engine."""
import csv
import io

import pytest

from aerotrace import uploads
from aerotrace.engine import AeroTrace
from aerotrace.ingest import FAA_COLUMNS
from aerotrace.uploads import UploadError, UploadStore, parse, validate

PAGE1 = """Airworthiness Directives; Zephyr Aero Helicopters
Docket No. FAA-2099-0001
SUMMARY: The FAA proposes to adopt a new airworthiness directive (AD) for certain Zephyr Aero Model ZX-100
helicopters. This proposed AD would require inspecting the tail rotor pitch link for cracks.
(a) Comments Due Date
The FAA must receive comments on this proposed AD by December 1, 2099.
(b) Affected ADs
None.
(c) Applicability
This AD applies to Zephyr Aero Model ZX-100 helicopters, serial numbers 0100"""
PAGE2 = """through 0199, certificated in any category.
(d) Subject
Joint Aircraft System Component (JASC) Code 6720, Tail Rotor Control System.
(e) Required Actions
Within 50 hours time-in-service, inspect the tail rotor pitch link for cracks and replace any cracked link."""


def make_pdf(pages: list[str]) -> bytes:
    """Smallest valid text PDF: one Helvetica text line per input line."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        lines = [ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for ln in text.split("\n")]
        stream = "BT /F1 10 Tf 12 TL 40 800 Td " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents {len(objs)} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Airworthiness Directives; Zephyr Aero Helicopters", level=1)
    d.add_paragraph("Docket No. FAA-2099-0002")
    d.add_heading("Applicability", level=2)
    d.add_paragraph("This AD applies to Zephyr Aero Model ZX-200 helicopters, serial numbers 0500 through 0599.")
    d.add_heading("Required Actions", level=2)
    d.add_paragraph("Replace the main rotor damper within 100 hours time-in-service.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Part number", "Action"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "ZD-4410", "Replace"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def official_rows(n=1):
    with open("data/ad_nprm.tsv", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter="\t")
        return reader.fieldnames, [next(reader) for _ in range(n)]


# ---------- validation ----------

@pytest.mark.parametrize("name,data,msg", [
    ("notes.txt", b"hello", "not supported"),
    ("old.doc", b"\xd0\xcf\x11\xe0", "save it as .docx"),
    ("old.xls", b"\xd0\xcf\x11\xe0", "save it as .xlsx"),
    ("data.csv", b"a,b\n1,2", "save it as .xlsx or tab-separated .tsv"),
    ("virus.exe", b"MZ", "not supported"),
    ("noext", b"x", "no extension"),
    ("fake.pdf", b"MZ\x90\x00 not a pdf", "not a PDF"),
    ("fake.docx", b"plain text", "not a valid Word file"),
    ("fake.xlsx", make_docx(), "not a Excel document"),
    ("empty.pdf", b"", "empty"),
    ("commas.tsv", b"a,b,c\n1,2,3\n", "no tab characters"),
])
def test_rejects(name, data, msg):
    with pytest.raises(UploadError, match=msg):
        validate(name, data)


def test_size_limit(monkeypatch):
    monkeypatch.setattr(uploads, "MAX_BYTES", 10)
    with pytest.raises(UploadError, match="limit"):
        validate("big.pdf", b"%PDF-" + b"x" * 20)


def test_accepts_all_four_kinds():
    assert validate("a.pdf", make_pdf(["hello"])) == "PDF"
    assert validate("a.DOCX", make_docx()) == "Word"
    assert validate("a.xlsx", make_xlsx({"S": [["a"], ["b"]]})) == "Excel"
    assert validate("a.tsv", b"a\tb\n1\t2\n") == "TSV"


def test_scanned_pdf_rejected():
    with pytest.raises(UploadError, match="No readable text"):
        parse("scan.pdf", make_pdf([""]))


# ---------- parsing ----------

def test_pdf_sections_across_pages():
    r = parse("zephyr.pdf", make_pdf([PAGE1, PAGE2]))
    doc = r["docs"][0]
    assert doc["docket"] == "FAA-2099-0001"
    assert doc["title"].startswith("Airworthiness Directives; Zephyr Aero")
    assert "ZX-100" in doc["models"]
    assert doc["upload"]["kind"] == "PDF" and "user-provided" in doc["source_kind"]
    by_section = {p["section"]: p for p in r["passages"]}
    applic = by_section["Applicability"]
    assert "0100 through 0199" in applic["text"]  # joined across the page break
    assert applic["page"] == 1 and by_section["Required Actions"]["page"] == 2


def test_docx_headings_and_tables():
    r = parse("zephyr.docx", make_docx())
    doc = r["docs"][0]
    assert doc["docket"] == "FAA-2099-0002" and "ZX-200" in doc["models"]
    sections = {p["section"] for p in r["passages"]}
    assert {"Applicability", "Required Actions"} <= sections
    table = [p for p in r["passages"] if p["paragraph"].startswith("table")]
    assert table and "Part number: ZD-4410" in table[0]["text"]


def test_excel_faa_layout_and_generic_sheet():
    header, rows = official_rows(1)
    data = make_xlsx({"NPRMs": [header, [rows[0][h] for h in header]],
                      "Fleet": [["Tail", "Model", "Serial"], [None, None, None], ["N123", "R22", "1200"]]})
    r = parse("mixed.xlsx", data)
    faa = [d for d in r["docs"] if d["doc_id"].startswith("upload/mixed/NPRMs/")]
    assert faa and faa[0]["docket"] == rows[0]["Docket Number"].strip()
    fleet = next(d for d in r["docs"] if d["doc_id"] == "upload/mixed/Fleet")
    assert fleet["models"] == ["R22"]
    row = next(p for p in r["passages"] if p["doc_id"] == fleet["doc_id"] and p["paragraph"] != "metadata")
    assert row["page"] == "Fleet row 3" and "Serial: 1200" in row["text"]  # real row number, blank row skipped


def test_tsv_generic():
    r = parse("parts.tsv", b"Part\tLimit\nZD-4410\t500 hours\n")
    assert any("Limit: 500 hours" in p["text"] for p in r["passages"])


def test_faa_columns_match_official_tsv():
    header, _ = official_rows(1)
    assert FAA_COLUMNS <= set(header)


# ---------- store + engine ----------

def test_store_dedups_and_deletes(tmp_path):
    store = UploadStore(tmp_path)
    data = make_pdf([PAGE1, PAGE2])
    p1 = store.save("a.pdf", data, {"x": {"title": "T"}})
    assert store.save("renamed.pdf", data) == p1  # same bytes, stored once
    assert [i["file"] for i in store.items()] == ["a.pdf"]
    assert store.delete(uploads.file_hash(data)) and store.items() == []


@pytest.fixture(scope="module")
def engine(tmp_path_factory):
    return AeroTrace("data/ad_nprm.tsv", "data/pdfs", tmp_path_factory.mktemp("uploads"))


def test_engine_upload_ask_remove(engine):
    data = make_pdf([PAGE1, PAGE2])
    q = "Does FAA-2099-0001 apply to Model ZX-100 serial number 0150?"
    assert engine.ask(q)["reason_code"] == "OUT_OF_SCOPE"

    out = engine.add_upload("zephyr.pdf", data)
    assert out["status"] == "added" and out["docs"] == ["upload/zephyr"]
    r = engine.ask(q)
    assert r["status"] == "ANSWERED" and r["applicability"]["result"] == "APPLIES"
    assert r["statements"][0]["citations"][0]["uploaded"] is True
    assert engine.ask(q.replace("0150", "0250"))["applicability"]["result"] == "DOES_NOT_APPLY"

    assert engine.add_upload("again.pdf", data)["status"] == "already_added"
    # survives a restart: a new engine on the same folder loads it from disk
    assert "upload/zephyr" in AeroTrace("data/ad_nprm.tsv", "data/pdfs", engine.uploads.folder).docs

    assert engine.remove_upload(uploads.file_hash(data)) == 1
    assert engine.ask(q)["reason_code"] == "OUT_OF_SCOPE"


def test_engine_review_overrides(engine):
    data = make_pdf([PAGE1.replace("FAA-2099-0001", "unknown"), PAGE2])
    out = engine.add_upload("nodocket.pdf", data, {"upload/nodocket": {"docket": "FAA-2099-0009", "models": ["ZX-100"]}})
    assert out["status"] == "added"
    assert engine.docs["upload/nodocket"]["docket"] == "FAA-2099-0009"
    stored = engine.uploads.items()[0]["overrides"]
    assert stored["upload/nodocket"]["docket"] == "FAA-2099-0009"  # reapplied on restart
    engine.remove_upload(uploads.file_hash(data))


def test_reuploading_official_rows_is_a_duplicate(engine):
    header, rows = official_rows(2)
    tsv = "\t".join(header) + "\n" + "\n".join("\t".join(r[h] for h in header) for r in rows) + "\n"
    out = engine.add_upload("copy.tsv", tsv.encode())
    assert out["status"] == "duplicate" and len(out["skipped"]) == 2
    assert engine.uploads.items() == []  # nothing stored
